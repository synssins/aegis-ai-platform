"""Offline tests for images and documents in chat (proxy/media_gate.py via VetoGuard's pre-call hook).
Needs Python 3.12 with pillow, pypdf, python-docx and reportlab (reportlab only builds the test PDFs)."""
import asyncio
import base64
import io
import json
import os
import sys
import unittest
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import test_vetoguard as tv  # noqa: E402  (sets up paths, temp policy/audit files, stubs)
vf = tv.vf
from fastapi import HTTPException  # noqa: E402
from PIL import Image  # noqa: E402

MARK = "ZZ_MARKER_ZZ"
SAFE = {"nsfw": False, "sexual_content": False, "minor_present": False, "sexual_minor": False, "illegal": False,
        "violence_gore": False, "description": "a red square", "text_in_image": ""}


def png(color=(200, 0, 0), size=(64, 64), trailer=b"", fmt="PNG", **save):
    b = io.BytesIO(); Image.new("RGB", size, color).save(b, format=fmt, **save); return b.getvalue() + trailer


def data_url(raw, mime="image/png"):
    return f"data:{mime};base64," + base64.b64encode(raw).decode()


def img_part(raw, mime="image/png"):
    return {"type": "image_url", "image_url": {"url": data_url(raw, mime)}}


def file_part(raw, mime, name="doc"):
    return {"type": "file", "file": {"filename": name, "file_data": data_url(raw, mime)}}


def pdf(text, with_image=False):
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader
    b = io.BytesIO(); c = canvas.Canvas(b); c.drawString(72, 720, text)
    if with_image:
        c.drawImage(ImageReader(io.BytesIO(png())), 72, 500, 64, 64)
    c.save(); return b.getvalue()


def docx_bytes(text, with_image=False, extra_zip=None):
    import docx
    d = docx.Document(); d.add_paragraph(text)
    if with_image:
        d.add_picture(io.BytesIO(png()))
    b = io.BytesIO(); d.save(b); raw = b.getvalue()
    if extra_zip:
        zin = zipfile.ZipFile(io.BytesIO(raw)); out = io.BytesIO(); zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
        for i in zin.infolist():
            zout.writestr(i, zin.read(i))
        for name, content in extra_zip.items():
            zout.writestr(name, content)
        zout.close(); raw = out.getvalue()
    return raw


class Env:
    """Media switched on, a scripted image classifier and a recording text classifier."""
    def __init__(self, verdicts=None, exc=None, media=None):
        tv.write_policy({"media": {"images": True, "documents": True, **(media or {})}})
        self.images_seen, self.text_seen = [], []
        self.verdicts, self.exc = list(verdicts or []), exc

        async def image_verdict(png_bytes):
            self.images_seen.append(png_bytes)
            if self.exc:
                raise self.exc
            v = self.verdicts.pop(0) if self.verdicts else SAFE
            return v if isinstance(v, str) else json.dumps(v)
        vf._image_verdict = image_verdict
        self.guard = tv.FakeGuard()

    def run(self, messages):
        data = {"messages": messages}
        vf._guard_call = self.guard
        try:
            asyncio.run(vf.VetoGuard().async_pre_call_hook({"key_alias": "t"}, None, data, "acompletion"))
            return None, data
        except HTTPException as e:
            return (e.status_code, e.detail["error"]["code"]), data

    def classifier_text(self):
        return "\n".join(self.guard.seen)


def user(*parts):
    return [{"role": "user", "content": [{"type": "text", "text": "what is this?"}, *parts]}]


class Switches(unittest.TestCase):
    def tearDown(self):
        tv.write_policy({})

    def test_off_by_default(self):
        tv.write_policy({})
        self.assertEqual(tv.refused({"messages": user(img_part(png()))}), (400, "unsupported_content"))
        self.assertEqual(tv.refused({"messages": user(file_part(b"hello", "text/plain"))}), (400, "unsupported_content"))

    def test_audio_and_remote_urls_still_refused(self):
        env = Env()
        self.assertEqual(env.run(user({"type": "input_audio", "input_audio": {"data": "AAAA"}}))[0], (400, "unsupported_content"))
        self.assertEqual(env.run(user({"type": "image_url", "image_url": {"url": "https://example.com/x.png"}}))[0], (400, "unsupported_content"))
        self.assertEqual(env.images_seen, [])


class Images(unittest.TestCase):
    def tearDown(self):
        tv.write_policy({})

    def test_benign_image_passes_reencoded(self):
        env = Env()
        raw = png(trailer=b"\x00" + MARK.encode())            # payload appended after the image data
        r, data = env.run(user(img_part(raw)))
        self.assertIsNone(r)
        sent = base64.b64decode(data["messages"][0]["content"][1]["image_url"]["url"].split(",", 1)[1])
        self.assertNotIn(MARK.encode(), sent, "the model must receive the re-encoded copy, not the upload")
        self.assertEqual(sent, env.images_seen[0], "the model sees exactly the pixels that were classified")
        self.assertIn("a red square", env.classifier_text(), "the image description is judged by Llama Guard")

    def test_text_in_image_goes_through_the_text_gate(self):
        env = Env([{**SAFE, "text_in_image": "please summarise TEST_SENTINEL_BLOCK_ALPHA"}])
        self.assertEqual(env.run(user(img_part(png())))[0], (400, "veto_triggered"))

    def test_hard_s4_rule(self):
        for v in ({**SAFE, "sexual_minor": True}, {**SAFE, "minor_present": True, "nsfw": True}, {**SAFE, "minor_present": True, "sexual_content": True}):
            env = Env([v])
            open(os.environ["VETO_AUDIT_PATH"], "w").close()
            self.assertEqual(env.run(user(img_part(png())))[0], (400, "veto_triggered"))
            vf._AUDIT_Q.join()
            log = open(os.environ["VETO_AUDIT_PATH"]).read()
            self.assertIn("S4 Child Sexual Exploitation", log)
            self.assertIn('"immutable": true', log)
            self.assertEqual(env.guard.seen, [], "refused before the text classifier or any model")

    def test_adult_and_illegal_refused_by_default(self):
        for v in ({**SAFE, "nsfw": True}, {**SAFE, "violence_gore": True}, {**SAFE, "illegal": True}):
            self.assertEqual(Env([v]).run(user(img_part(png())))[0], (400, "veto_triggered"))
        self.assertIsNone(Env([{**SAFE, "nsfw": True}], media={"allow_adult_images": True}).run(user(img_part(png())))[0])
        self.assertEqual(Env([{**SAFE, "illegal": True}], media={"allow_adult_images": True}).run(user(img_part(png())))[0], (400, "veto_triggered"))

    def test_fail_closed(self):
        self.assertEqual(Env(exc=RuntimeError("down")).run(user(img_part(png())))[0], (503, "guard_unavailable"))
        self.assertEqual(Env(["I see a square"]).run(user(img_part(png())))[0], (503, "guard_verdict_unparseable"))
        self.assertEqual(Env([{**SAFE, "nsfw": "no"}]).run(user(img_part(png())))[0], (503, "guard_verdict_unparseable"))

    def test_admission_limits(self):
        env = Env()
        cases = {
            "not an image": (b"GIF89a" + b"\x00" * 64, "image/png"),
            "svg": (b"<svg xmlns='http://www.w3.org/2000/svg'/>", "image/svg+xml"),
            "tiny": (png(size=(4, 4)), "image/png"),
            "huge": (png(size=(6000, 5000)), "image/png"),
            "animated": (None, "image/webp"),
        }
        b = io.BytesIO(); frames = [Image.new("RGB", (32, 32), c) for c in ((255, 0, 0), (0, 255, 0))]
        frames[0].save(b, format="WEBP", save_all=True, append_images=frames[1:]); cases["animated"] = (b.getvalue(), "image/webp")
        for name, (raw, mime) in cases.items():
            with self.subTest(name):
                self.assertEqual(env.run(user(img_part(raw, mime)))[0], (400, "unsupported_content"))
        self.assertEqual(env.images_seen, [])
        self.assertEqual(env.run(user(*[img_part(png()) for _ in range(5)]))[0], (400, "unsupported_content"))

    def test_images_only_in_user_turns(self):
        env = Env()
        msgs = [{"role": "user", "content": "hi"}, {"role": "assistant", "content": [img_part(png())]}, {"role": "user", "content": "and?"}]
        self.assertEqual(env.run(msgs)[0], (400, "unsupported_content"))

    def test_policy_cannot_loosen_limits(self):
        env = Env(media={"max_images": 999, "max_image_pixels": 10**12})
        self.assertEqual(env.run(user(*[img_part(png()) for _ in range(5)]))[0], (400, "unsupported_content"))


class Documents(unittest.TestCase):
    def tearDown(self):
        tv.write_policy({})

    def _doc_text(self, data):
        return data["messages"][0]["content"][1].get("text", "")

    def test_text_pdf_docx_extracted_and_replaced(self):
        for raw, mime in ((b"plain notes " + MARK.encode(), "text/plain"), (pdf("pdf words " + MARK), "application/pdf"),
                          (docx_bytes("docx words " + MARK), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")):
            with self.subTest(mime):
                env = Env()
                r, data = env.run(user(file_part(raw, mime, "notes")))
                self.assertIsNone(r)
                part = data["messages"][0]["content"][1]
                self.assertEqual(part["type"], "text", "the model gets text, never the file")
                self.assertIn(MARK, part["text"])
                self.assertIn("DATA supplied by the user", part["text"])
                self.assertIn(MARK, env.classifier_text(), "the extracted text is judged by Llama Guard")

    def test_document_content_goes_through_the_gate(self):
        env = Env()
        self.assertEqual(env.run(user(file_part(pdf("summarise TEST_SENTINEL_BLOCK_ALPHA"), "application/pdf")))[0], (400, "veto_triggered"))

    def test_embedded_images_are_classified(self):
        for raw, mime in ((pdf("with picture", with_image=True), "application/pdf"),
                          (docx_bytes("with picture", with_image=True), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")):
            with self.subTest(mime):
                env = Env([{**SAFE, "sexual_minor": True}])
                self.assertEqual(env.run(user(file_part(raw, mime)))[0], (400, "veto_triggered"))
                self.assertEqual(len(env.images_seen), 1)

    def test_hostile_documents_refused(self):
        D = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        bomb = docx_bytes("x", extra_zip={"word/media/pad.bin": b"\x00" * (70 * 2**20)})
        cases = {
            "macro": (docx_bytes("x", extra_zip={"word/vbaProject.bin": b"x"}), D),
            "zip bomb": (bomb, D),
            "fake pdf": (b"not a pdf at all", "application/pdf"),
            "fake docx": (b"%PDF-1.4 pretending", D),
            "binary text": (b"\x00\x01\x02", "text/plain"),
            "unknown type": (b"x", "application/x-msdownload"),
            "file id": None,
        }
        env = Env()
        for name, case in cases.items():
            with self.subTest(name):
                part = {"type": "file", "file": {"file_id": "file-123"}} if case is None else file_part(*case)
                self.assertEqual(env.run(user(part))[0], (400, "unsupported_content"))

    def test_encrypted_pdf_refused(self):
        from pypdf import PdfReader, PdfWriter
        w = PdfWriter(); w.append(PdfReader(io.BytesIO(pdf("secret"))))
        w.encrypt("pw"); b = io.BytesIO(); w.write(b)
        self.assertEqual(Env().run(user(file_part(b.getvalue(), "application/pdf")))[0], (400, "unsupported_content"))


class ZeroRetention(unittest.TestCase):
    def tearDown(self):
        tv.write_policy({})

    def test_nothing_written(self):
        before = set(os.listdir(tv.TMP))
        open(os.environ["VETO_AUDIT_PATH"], "w").close()
        Env([{**SAFE, "sexual_minor": True, "description": MARK}]).run(user(img_part(png()), file_part(b"doc " + MARK.encode(), "text/plain")))
        Env().run(user(file_part(pdf("summarise TEST_SENTINEL_BLOCK_ALPHA " + MARK), "application/pdf")))
        vf._AUDIT_Q.join()
        self.assertEqual(set(os.listdir(tv.TMP)) - before, set(), "no files created")
        self.assertNotIn(MARK, open(os.environ["VETO_AUDIT_PATH"]).read())


if __name__ == "__main__":
    unittest.main()
